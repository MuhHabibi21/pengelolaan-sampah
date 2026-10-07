import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Colors
    C_EMERALD = RGBColor(5, 150, 105)      # #059669
    C_DARK = RGBColor(6, 78, 59)           # #064E3B
    C_TEAL = RGBColor(13, 148, 136)        # #0D9488
    C_BG = RGBColor(248, 250, 252)         # #F8FAFC
    C_WHITE = RGBColor(255, 255, 255)
    C_SLATE = RGBColor(30, 41, 59)         # #1E293B
    C_MUTED = RGBColor(100, 116, 139)      # #64748B
    C_ACCENT = RGBColor(16, 185, 129)      # #10B981
    C_BOX_BORDER = RGBColor(203, 213, 225)
    C_PLACEHOLDER_BG = RGBColor(241, 245, 249)

    def set_slide_bg(slide, color=C_BG):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, tag, title, subtitle=None):
        # Tag pill
        tag_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(2.2), Inches(0.35))
        tag_box.fill.solid()
        tag_box.fill.fore_color.rgb = RGBColor(209, 250, 229)
        tag_box.line.fill.background()
        tf = tag_box.text_frame
        tf.text = tag.upper()
        p = tf.paragraphs[0]
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_DARK
        p.alignment = PP_ALIGN.CENTER

        # Title
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.5), Inches(0.6))
        tf = tx_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = C_SLATE

        if subtitle:
            p2 = tf.add_paragraph()
            p2.text = subtitle
            p2.font.size = Pt(11)
            p2.font.color.rgb = C_MUTED

    def add_card(slide, x, y, w, h, bg_color=C_WHITE, border_color=C_BOX_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)
        return card

    def add_screenshot_box(slide, x, y, w, h, title="TEMPATKAN SCREENSHOT DI SINI", subtitle="Klik ikon gambar / paste screenshot"):
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        box.fill.solid()
        box.fill.fore_color.rgb = C_PLACEHOLDER_BG
        box.line.color.rgb = C_EMERALD
        box.line.width = Pt(1.5)
        
        tf = box.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = f"📷 {title}"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = C_DARK
        p.alignment = PP_ALIGN.CENTER
        
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = C_MUTED
        p2.alignment = PP_ALIGN.CENTER
        return box

    # -------------------------------------------------------------
    # SLIDE 1: COVER
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1, C_DARK)

    # Accent decorative block
    dec = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.5), 0, Inches(4.833), prs.slide_height)
    dec.fill.solid()
    dec.fill.fore_color.rgb = C_EMERALD
    dec.line.fill.background()

    # Cover Text Frame
    tb = s1.shapes.add_textbox(Inches(0.9), Inches(1.5), Inches(7.2), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "BUKU PANDUAN PENGGUNA • USER MANUAL"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = RGBColor(167, 243, 208)

    p1 = tf.add_paragraph()
    p1.text = "PEDULI SAMPAH"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = C_WHITE

    p2 = tf.add_paragraph()
    p2.text = "Sistem Informasi Pengelolaan & Bank Sampah Berbasis Web Terpadu"
    p2.font.size = Pt(16)
    p2.font.color.rgb = RGBColor(226, 232, 240)

    p3 = tf.add_paragraph()
    p3.text = "\n\nDisusun Oleh:\nMuhammad Habibi El Islami\nTugas Ujian Tengah Semester (UTS) Pemrograman Web\nNext.js 15 • PostgreSQL • Prisma ORM • Tailwind CSS"
    p3.font.size = Pt(12)
    p3.font.color.rgb = RGBColor(203, 213, 225)

    # Right side decorative stats
    rtb = s1.shapes.add_textbox(Inches(8.8), Inches(1.8), Inches(4.0), Inches(4.0))
    rtf = rtb.text_frame
    rtf.word_wrap = True
    rp0 = rtf.paragraphs[0]
    rp0.text = "EDISI 2026"
    rp0.font.size = Pt(36)
    rp0.font.bold = True
    rp0.font.color.rgb = C_WHITE

    rp1 = rtf.add_paragraph()
    rp1.text = "\n✓ Role Warga & Admin\n✓ Transaksi Bank Sampah\n✓ Laporan & Foto Autentik\n✓ Relasi 1:1, 1:N, N:N\n✓ Deployment Live Vercel"
    rp1.font.size = Pt(14)
    rp1.font.color.rgb = RGBColor(236, 253, 245)

    # -------------------------------------------------------------
    # SLIDE 2: DAFTAR ISI & TUJUAN
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2)
    add_header(s2, "Overview Dokumen", "Daftar Isi & Tujuan Panduan", "Struktur panduan lengkap untuk kemudahan navigasi pengguna.")

    add_card(s2, 0.8, 1.8, 5.7, 5.0)
    tb = s2.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.3), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "📑 DAFTAR KONTEN UTAMA"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_DARK

    items = [
        "1. Pendahuluan & Arsitektur Sistem",
        "2. Matriks Hak Akses (Role Warga vs Admin)",
        "3. Panduan Pengguna: Registrasi & Login",
        "4. Panduan Warga: Laporan Pengaduan + Bukti Foto",
        "5. Panduan Warga: Transaksi Setor Bank Sampah",
        "6. Panduan Admin: Validasi Setoran & Pengguna",
        "7. Panduan Admin: Master Data & Fasilitas (N:N)",
        "8. Relasi Database (ERD) & Panduan Troubleshooting",
    ]
    for it in items:
        p = tf.add_paragraph()
        p.text = f"• {it}"
        p.font.size = Pt(11.5)
        p.font.color.rgb = C_SLATE

    add_card(s2, 6.8, 1.8, 5.7, 5.0)
    tb2 = s2.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.6))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "🎯 TUJUAN DAN SASARAN APLIKASI"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_DARK

    p = tf2.add_paragraph()
    p.text = "\n1. Digitalisasi Bank Sampah:\nMemfasilitasi warga mencatat setoran sampah daur ulang dan menerima reward berupa saldo & poin."
    p.font.size = Pt(11.5)
    p.font.color.rgb = C_SLATE

    p = tf2.add_paragraph()
    p.text = "\n2. Transparansi Pengaduan Publik:\nMenyediakan kanal pelaporan tumpukan sampah liar dengan bukti foto dan data berat akurat."
    p.font.size = Pt(11.5)
    p.font.color.rgb = C_SLATE

    p = tf2.add_paragraph()
    p.text = "\n3. Efisiensi Pengawasan Kota:\nAdmin dapat memantau perputaran sampah, verifikasi setor, dan persebaran fasilitas secara terpadu."
    p.font.size = Pt(11.5)
    p.font.color.rgb = C_SLATE

    # -------------------------------------------------------------
    # SLIDE 3: MATRIKS HAK AKSES
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3)
    add_header(s3, "Hak Akses Pengguna", "Matriks Peran: Warga (User) vs Administrator", "Pembagian fitur dan wewenang operasional pada sistem Peduli Sampah.")

    add_card(s3, 0.8, 1.8, 5.7, 5.0)
    tb = s3.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.3), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "👤 ROLE WARGA / MASYARAKAT"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = C_DARK

    warga_features = [
        "✓ Registrasi Akun Warga Baru mandiri",
        "✓ Login Akun Warga dengan JWT Cookie",
        "✓ Dashboard Ringkasan Laporan, Berat, & Poin",
        "✓ Membuat Pengaduan Sampah + Upload Foto (1:1)",
        "✓ Transaksi Setor Sampah + Estimasi Otomatis",
        "✓ Melihat Riwayat Laporan Pribadi",
        "✓ Mengubah Profil & Ganti Kata Sandi",
    ]
    for wf in warga_features:
        p = tf.add_paragraph()
        p.text = wf
        p.font.size = Pt(11)
        p.font.color.rgb = C_SLATE

    add_card(s3, 6.8, 1.8, 5.7, 5.0)
    tb = s3.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🛡️ ROLE ADMINISTRATOR PENGELOLA"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = C_DARK

    admin_features = [
        "✓ Dashboard Ikhtisar Statistik Sampah Kota",
        "✓ Validasi & Verifikasi Transaksi (Selesai/Tolak)",
        "✓ Monitoring & Pengawasan Seluruh Laporan Masuk",
        "✓ Kelola Pengguna (Ubah Role Warga ⇄ Admin)",
        "✓ Kelola Master Kategori Sampah (CRUD)",
        "✓ Kelola Master Data Wilayah Kecamatan (CRUD)",
        "✓ Kelola Fasilitas Daur Ulang (Relasi N:N)",
    ]
    for af in admin_features:
        p = tf.add_paragraph()
        p.text = af
        p.font.size = Pt(11)
        p.font.color.rgb = C_SLATE

    # -------------------------------------------------------------
    # SLIDE 4: WARGA - REGISTRASI & LOGIN
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4)
    add_header(s4, "Panduan Warga", "1. Registrasi Akun Baru & Autentikasi Login", "Tata cara pendaftaran akun warga dan masuk ke sistem.")

    # Left: Registrasi
    add_screenshot_box(s4, 0.8, 1.8, 5.7, 3.4, "Gambar 1: Halaman Registrasi (/register)")
    add_card(s4, 0.8, 5.35, 5.7, 1.6)
    tb = s4.shapes.add_textbox(Inches(0.95), Inches(5.45), Inches(5.4), Inches(1.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Langkah Registrasi:"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = "1. Buka menu 'Daftar Akun Baru' di pojok kanan atas.\n2. Lengkapi Nama Lengkap, Email, No HP, 16 Digit NIK, dan Password.\n3. Klik 'Daftar Akun Warga' untuk langsung diarahkan ke Dashboard."
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_SLATE

    # Right: Login
    add_screenshot_box(s4, 6.8, 1.8, 5.7, 3.4, "Gambar 2: Halaman Login (/login)")
    add_card(s4, 6.8, 5.35, 5.7, 1.6)
    tb = s4.shapes.add_textbox(Inches(6.95), Inches(5.45), Inches(5.4), Inches(1.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Langkah Login:"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = "1. Kunjungi menu Masuk (/login).\n2. Masukkan alamat email terdaftar dan kata sandi.\n3. Akun Demo Warga: user@example.com | Kata Sandi: password123."
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_SLATE

    # -------------------------------------------------------------
    # SLIDE 5: WARGA - BERANDA & DASHBOARD
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s5)
    add_header(s5, "Panduan Warga", "2. Beranda Publik & Dashboard Kontribusi Warga", "Monitoring laporan terkini dan ringkasan kontribusi lingkungan.")

    # Left: Landing Page
    add_screenshot_box(s5, 0.8, 1.8, 5.7, 3.4, "Gambar 3: Beranda Publik (/)", "Statistik Publik & Tabel Laporan Terkini")
    add_card(s5, 0.8, 5.35, 5.7, 1.6)
    tb = s5.shapes.add_textbox(Inches(0.95), Inches(5.45), Inches(5.4), Inches(1.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Fitur Beranda Publik:"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = "• Menampilkan ringkasan total sampah terkelola di seluruh kota.\n• Tabel keterbukaan informasi publik dengan sensor privasi nama (Bud***).\n• Akses cepat untuk warga mulai melapor."
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_SLATE

    # Right: Dashboard
    add_screenshot_box(s5, 6.8, 1.8, 5.7, 3.4, "Gambar 4: Dashboard Warga (/dashboard)", "Kartu Statistik & Aksi Cepat")
    add_card(s5, 6.8, 5.35, 5.7, 1.6)
    tb = s5.shapes.add_textbox(Inches(6.95), Inches(5.45), Inches(5.4), Inches(1.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Fitur Dashboard Warga:"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = "• Kartu Total Laporan Saya, Akumulasi Berat Sampah, & Poin Reward.\n• Tombol pintas untuk langsung membuat pengaduan atau setor sampah.\n• Informasi sambutan personal warga."
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_SLATE

    # -------------------------------------------------------------
    # SLIDE 6: WARGA - LAPORAN & UPLOAD FOTO
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s6)
    add_header(s6, "Panduan Warga", "3. Kirim Pengaduan Sampah (+ Upload Foto 1:1)", "Formulir pelaporan tumpukan sampah dan riwayat pengaduan.")

    add_screenshot_box(s6, 0.8, 1.8, 5.7, 3.4, "Gambar 5: Form Laporan Baru (/dashboard/laporan/baru)")
    add_card(s6, 0.8, 5.35, 5.7, 1.6)
    tb = s6.shapes.add_textbox(Inches(0.95), Inches(5.45), Inches(5.4), Inches(1.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Alur Pembuatan Laporan (1-to-1):"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = "1. Pilih Jenis Sampah & Wilayah Kecamatan.\n2. Masukkan Berat Sampah (Wajib > 0 Kg, divalidasi Zod).\n3. Lampirkan file foto autentik lalu klik 'Kirim Laporan'."
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_SLATE

    add_screenshot_box(s6, 6.8, 1.8, 5.7, 3.4, "Gambar 6: Riwayat Laporan Warga (/dashboard/laporan)")
    add_card(s6, 6.8, 5.35, 5.7, 1.6)
    tb = s6.shapes.add_textbox(Inches(6.95), Inches(5.45), Inches(5.4), Inches(1.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Fitur Riwayat Laporan:"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = "• Memantau seluruh histori laporan yang pernah diajukan.\n• Dilengkapi tanggal lengkap, status wilayah, dan thumbnail foto bukti.\n• Memudahkan verifikasi data saat petugas terjun ke lokasi."
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_SLATE

    # -------------------------------------------------------------
    # SLIDE 7: WARGA - TRANSAKSI & SETTINGS
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s7)
    add_header(s7, "Panduan Warga", "4. Transaksi Setor Bank Sampah & Pengaturan Akun", "Kalkulasi otomatis nilai rupiah/poin dan manajemen profil.")

    add_screenshot_box(s7, 0.8, 1.8, 5.7, 3.4, "Gambar 7: Transaksi Setor Sampah (/dashboard/transaksi)")
    add_card(s7, 0.8, 5.35, 5.7, 1.6)
    tb = s7.shapes.add_textbox(Inches(0.95), Inches(5.45), Inches(5.4), Inches(1.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Kalkulasi Transaksi Otomatis:"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = "• Memilih jenis sampah & memasukkan berat langsung mengkalkulasi harga.\n• Sistem menghitung total poin reward (10 poin/kg).\n• Transaksi berstatus PENDING hingga divalidasi oleh petugas admin."
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_SLATE

    add_screenshot_box(s7, 6.8, 1.8, 5.7, 3.4, "Gambar 8: Pengaturan Akun (/dashboard/settings)")
    add_card(s7, 6.8, 5.35, 5.7, 1.6)
    tb = s7.shapes.add_textbox(Inches(6.95), Inches(5.45), Inches(5.4), Inches(1.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Fitur Pengaturan Profil:"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = "• Ubah data profil nama lengkap dan nomor HP aktif.\n• Ganti kata sandi akun dengan enkripsi bcrypt yang aman.\n• Menampilkan status akun warga terverifikasi."
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_SLATE

    # -------------------------------------------------------------
    # SLIDE 8: ADMIN - DASHBOARD & TRANSAKSI
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s8)
    add_header(s8, "Panduan Admin", "1. Dashboard Panel Admin & Validasi Transaksi", "Pusat komando pengawasan dan persetujuan setoran bank sampah.")

    add_screenshot_box(s8, 0.8, 1.8, 5.7, 3.4, "Gambar 9: Dashboard Admin (/admin)")
    add_card(s8, 0.8, 5.35, 5.7, 1.6)
    tb = s8.shapes.add_textbox(Inches(0.95), Inches(5.45), Inches(5.4), Inches(1.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Ikhtisar Dashboard Admin:"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = "• Ringkasan 4 metrik utama: Total Laporan, Berat (Kg), Transaksi, & Warga.\n• Kartu antrean transaksi setor yang menunggu validasi petugas.\n• Tautan pintas ke modul kelola operasional."
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_SLATE

    add_screenshot_box(s8, 6.8, 1.8, 5.7, 3.4, "Gambar 10: Kelola Transaksi (/admin/transaksi)")
    add_card(s8, 6.8, 5.35, 5.7, 1.6)
    tb = s8.shapes.add_textbox(Inches(6.95), Inches(5.45), Inches(5.4), Inches(1.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Alur Verifikasi Transaksi:"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = "• Admin memeriksa kesesuaian sampah yang disetor oleh warga.\n• Tombol Selesai (Hijau): Menyetujui dan mencairkan poin/saldo warga.\n• Tombol Tolak (Merah): Membatalkan transaksi jika sampah tidak layak."
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_SLATE

    # -------------------------------------------------------------
    # SLIDE 9: ADMIN - LAPORAN & USER MANAGEMENT
    # -------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s9)
    add_header(s9, "Panduan Admin", "2. Manajemen Semua Laporan & Delegasi Pengguna", "Monitoring seluruh laporan sekota dan pengaturan wewenang hak akses.")

    add_screenshot_box(s9, 0.8, 1.8, 5.7, 3.4, "Gambar 11: Semua Laporan Masuk (/admin/laporan)")
    add_card(s9, 0.8, 5.35, 5.7, 1.6)
    tb = s9.shapes.add_textbox(Inches(0.95), Inches(5.45), Inches(5.4), Inches(1.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Monitoring Laporan Publik:"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = "• Admin dapat melihat semua laporan dari seluruh kecamatan.\n• Menampilkan identitas warga pelapor, waktu pelaporan, dan berat.\n• Dilengkapi foto bukti asli untuk verifikasi ke lapangan."
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_SLATE

    add_screenshot_box(s9, 6.8, 1.8, 5.7, 3.4, "Gambar 12: Kelola Pengguna (/admin/users)")
    add_card(s9, 6.8, 5.35, 5.7, 1.6)
    tb = s9.shapes.add_textbox(Inches(6.95), Inches(5.45), Inches(5.4), Inches(1.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Fitur Kelola Pengguna:"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = "• Menampilkan daftar seluruh akun warga dan petugas.\n• Fitur Switch Role: Klik 'Jadikan Admin' atau 'Jadikan Warga'.\n• Perubahan hak akses berlaku secara instan via Server Action."
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_SLATE

    # -------------------------------------------------------------
    # SLIDE 10: ADMIN - MASTER DATA & FASILITAS N:N
    # -------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s10)
    add_header(s10, "Panduan Admin", "3. Master Data Jenis, Wilayah & Fasilitas (N:N)", "Pengelolaan data master dan implementasi relasi Many-to-Many.")

    # 3 screenshot / cards
    add_screenshot_box(s10, 0.8, 1.8, 3.7, 3.4, "Gambar 13: Jenis Sampah", "/admin/jenis-sampah")
    add_card(s10, 0.8, 5.35, 3.7, 1.6)
    tb = s10.shapes.add_textbox(Inches(0.9), Inches(5.45), Inches(3.5), Inches(1.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Master Jenis Sampah:"
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = "• Tambah kategori sampah baru.\n• Hapus kategori yang tidak aktif."
    p.font.size = Pt(9)
    p.font.color.rgb = C_SLATE

    add_screenshot_box(s10, 4.8, 1.8, 3.7, 3.4, "Gambar 14: Master Wilayah", "/admin/wilayah")
    add_card(s10, 4.8, 5.35, 3.7, 1.6)
    tb = s10.shapes.add_textbox(Inches(4.9), Inches(5.45), Inches(3.5), Inches(1.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Master Wilayah:"
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = "• Tambah data kecamatan baru.\n• Pemetaan area jangkauan bank sampah."
    p.font.size = Pt(9)
    p.font.color.rgb = C_SLATE

    add_screenshot_box(s10, 8.8, 1.8, 3.7, 3.4, "Gambar 15: Fasilitas (N:N)", "/admin/fasilitas")
    add_card(s10, 8.8, 5.35, 3.7, 1.6)
    tb = s10.shapes.add_textbox(Inches(8.9), Inches(5.45), Inches(3.5), Inches(1.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Fasilitas (Many-to-Many):"
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = C_DARK
    p = tf.add_paragraph()
    p.text = "• Relasi N-to-N Fasilitas & Jenis.\n• Satu fasilitas menampung banyak jenis."
    p.font.size = Pt(9)
    p.font.color.rgb = C_SLATE

    # -------------------------------------------------------------
    # SLIDE 11: DATABASE ERD
    # -------------------------------------------------------------
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s11)
    add_header(s11, "Arsitektur Data", "Struktur Database Relasional & Relasi Prisma", "Implementasi relasi 1:1, 1:N, dan N:N sesuai spesifikasi tugas UTS.")

    add_card(s11, 0.8, 1.8, 5.7, 5.0)
    tb = s11.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.3), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "📊 DAFTAR RELASI ENTITAS (ERD)"
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = C_DARK

    p = tf.add_paragraph()
    p.text = "\n1. Relasi One-to-One (1-to-1):\n• LaporanSampah ↔ FotoSampah (Tiap laporan terikat 1 foto bukti via FK laporanId unik)."
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_SLATE

    p = tf.add_paragraph()
    p.text = "\n2. Relasi One-to-Many (1-to-N):\n• User → LaporanSampah (1 user memiliki banyak laporan).\n• User → TransaksiSampah (1 user memiliki banyak transaksi setor).\n• JenisSampah & Wilayah → Laporan & Transaksi."
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_SLATE

    p = tf.add_paragraph()
    p.text = "\n3. Relasi Many-to-Many (N-to-N):\n• Fasilitas ↔ JenisSampah (Fasilitas menampung banyak jenis, jenis sampah diproses di banyak fasilitas)."
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_SLATE

    add_card(s11, 6.8, 1.8, 5.7, 5.0)
    tb = s11.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🛡️ FITUR KEAMANAN & INTEGRITAS DATA"
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = C_DARK

    p = tf.add_paragraph()
    p.text = "\n• Enkripsi Kata Sandi: Menggunakan algoritma Bcrypt (10 salt rounds) untuk proteksi kredensial."
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_SLATE

    p = tf.add_paragraph()
    p.text = "\n• JWT HTTP-Only Cookies: Sesi disimpan aman di browser mencegah serangan XSS."
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_SLATE

    p = tf.add_paragraph()
    p.text = "\n• Validasi Input Zod: Menjamin berat > 0 kg, NIK 16 digit, dan validitas format email."
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_SLATE

    p = tf.add_paragraph()
    p.text = "\n• Prisma Atomic Transactions: Memastikan pembuatan data laporan dan berkas foto tersimpan utuh."
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_SLATE

    # -------------------------------------------------------------
    # SLIDE 12: CLOSING & THANK YOU
    # -------------------------------------------------------------
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s12, C_DARK)

    # Accent decorative block
    dec = s12.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(5.8), prs.slide_width, Inches(1.7))
    dec.fill.solid()
    dec.fill.fore_color.rgb = C_EMERALD
    dec.line.fill.background()

    tb = s12.shapes.add_textbox(Inches(1.5), Inches(1.5), Inches(10.3), Inches(4.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "TERIMA KASIH"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.alignment = PP_ALIGN.CENTER

    p = tf.add_paragraph()
    p.text = "Sistem Informasi Pengelolaan & Bank Sampah (Peduli Sampah)\nDokumen Buku Panduan Pengguna (User Manual) UTS Pemrograman Web"
    p.font.size = Pt(16)
    p.font.color.rgb = RGBColor(209, 250, 229)
    p.alignment = PP_ALIGN.CENTER

    p = tf.add_paragraph()
    p.text = "\nPengembang: Muhammad Habibi El Islami\nGitHub: https://github.com/MuhHabibi21/pengelolaan-sampah\nLive Website: https://pengelolaan-sampah-beta.vercel.app"
    p.font.size = Pt(13)
    p.font.color.rgb = C_WHITE
    p.alignment = PP_ALIGN.CENTER

    out_file = "D:/Koding/pengelolaan-sampah/Buku_Panduan_Peduli_Sampah.pptx"
    prs.save(out_file)
    print(f"PowerPoint Presentation successfully created at: {out_file}")

if __name__ == "__main__":
    create_deck()
